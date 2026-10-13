#!/usr/bin/env python3
"""SalesDoc API V2 dan oylik sotuv ma'lumotlarini tortib, Excel hisobot tayyorlaydi.

Faqat O'QISH (get*) metodlari ishlatiladi — SalesDoc'dagi hech narsa o'zgarmaydi.

Muhit o'zgaruvchilari (environment secrets):
    SALESDOC_DOMAIN    masalan: kompaniya.salesdoc.io
    SALESDOC_LOGIN     API uchun alohida foydalanuvchi logini
    SALESDOC_PASSWORD  shu foydalanuvchi paroli
    SALESDOC_FILIAL    (ixtiyoriy) filial kodi; filial bo'lmasa qo'ymang

Misollar:
    python3 salesdoc/salesdoc_export.py                       # shu oy, 1-sanadan bugungacha
    python3 salesdoc/salesdoc_export.py --product "happy toy" # faqat Happy Toy mahsulotlari
    python3 salesdoc/salesdoc_export.py --from 2026-09-01 --to 2026-09-30
"""

import argparse
import datetime as dt
import json
import os
import sys
import time
import urllib.error
import urllib.request
from collections import defaultdict

# Skript faqat shu metodlarni chaqira oladi: tasodifan set*/delete* yuborilmasligi uchun.
READ_ONLY_METHODS = {
    "login",
    "getFilials",
    "getAgent",
    "getClient",
    "getProduct",
    "getOrder",
    "getOrderDefect",
}

STATUS_NAMES = {1: "Yangi", 2: "Yuborilgan", 3: "Yetkazilgan", 4: "Yopilgan", 5: "Bekor qilingan"}
DONE_STATUSES = {3, 4}      # haqiqatda sotilgan
PENDING_STATUSES = {1, 2}   # hali yo'lda / tasdiqlanmagan
PAGE_LIMIT = 1000


class SalesDocError(RuntimeError):
    pass


class SalesDocClient:
    def __init__(self, domain, login, password, filial=None, timeout=60):
        domain = domain.strip().rstrip("/")
        if not domain.startswith("http"):
            domain = "https://" + domain
        self.url = domain + "/api/v2"
        self.login_name = login
        self.password = password
        self.filial = filial
        self.timeout = timeout
        self.auth = None

    def _post(self, body):
        if body["method"] not in READ_ONLY_METHODS:
            raise SalesDocError(f"Ruxsat etilmagan metod: {body['method']}")
        data = json.dumps(body).encode("utf-8")
        for attempt in range(5):
            req = urllib.request.Request(
                self.url, data=data, headers={"Content-Type": "application/json"}, method="POST"
            )
            try:
                with urllib.request.urlopen(req, timeout=self.timeout) as resp:
                    payload = json.loads(resp.read().decode("utf-8"))
            except urllib.error.HTTPError as e:
                if e.code == 429 and attempt < 4:
                    time.sleep(5 * (attempt + 1))
                    continue
                raise SalesDocError(f"{body['method']}: HTTP {e.code} — {e.read()[:300]!r}") from e
            except urllib.error.URLError as e:
                if attempt < 4:
                    time.sleep(2 ** (attempt + 1))
                    continue
                raise SalesDocError(f"{body['method']}: ulanib bo'lmadi — {e.reason}") from e

            if payload.get("status") is True:
                return payload
            err = payload.get("error") or {}
            if err.get("code") == 429 and attempt < 4:
                time.sleep(5 * (attempt + 1))
                continue
            raise SalesDocError(f"{body['method']}: {err.get('code')} {err.get('message')}")
        raise SalesDocError(f"{body['method']}: urinishlar tugadi")

    def login(self):
        payload = self._post(
            {"method": "login", "auth": {"login": self.login_name, "password": self.password}}
        )
        result = payload["result"]
        self.auth = {"userId": result["userId"], "token": result["token"]}

    def get_all(self, method, result_key, params=None):
        """Barcha sahifalarni yig'ib, bitta ro'yxat qaytaradi."""
        rows, page = [], 1
        while True:
            body = {
                "method": method,
                "auth": self.auth,
                "params": {**(params or {}), "limit": PAGE_LIMIT, "page": page},
            }
            if self.filial:
                body["filial"] = {"filial_id": self.filial}
            payload = self._post(body)
            batch = (payload.get("result") or {}).get(result_key) or []
            rows.extend(batch)
            pag = payload.get("pagination") or {}
            total = int(pag.get("total") or 0)
            if not batch or len(rows) >= total:
                return rows
            page += 1


def ref_key(obj):
    """Bitta ob'ekt uchun barqaror kalit: SD_id > CS_id > code_1C."""
    obj = obj or {}
    return obj.get("SD_id") or obj.get("CS_id") or obj.get("code_1C") or ""


def name_map(items):
    out = {}
    for it in items:
        for k in ("SD_id", "CS_id", "code_1C"):
            if it.get(k):
                out[it[k]] = it.get("name") or ""
    return out


def lookup(names, obj):
    obj = obj or {}
    for k in ("SD_id", "CS_id", "code_1C"):
        if obj.get(k) and obj[k] in names:
            return names[obj[k]]
    return ""


def num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def fetch(client, date_from, date_to):
    period = {"date": {"from": date_from, "to": date_to}}
    print("→ Agentlar...", file=sys.stderr)
    agents = client.get_all("getAgent", "agent")
    print("→ Mijozlar (do'konlar)...", file=sys.stderr)
    clients = client.get_all("getClient", "client")
    print("→ Mahsulotlar...", file=sys.stderr)
    products = client.get_all("getProduct", "product")
    print(f"→ Buyurtmalar {date_from} — {date_to}...", file=sys.stderr)
    orders = client.get_all(
        "getOrder",
        "order",
        {"filter": {"agent": "all", "include": "all", "status": [1, 2, 3, 4, 5], "period": period}},
    )
    print("→ Qaytarishlar...", file=sys.stderr)
    defects = client.get_all(
        "getOrderDefect", "order", {"filter": {"status": [1, 2, 3, 4, 5], "period": period}}
    )
    return agents, clients, products, orders, defects


def build_report(agents, clients, products, orders, defects, product_filter=None):
    agent_names = name_map(agents)
    client_names = name_map(clients)
    product_names = name_map(products)
    pf = [p.lower() for p in (product_filter or [])]

    def product_ok(prod):
        if not pf:
            return True
        name = (prod.get("name") or lookup(product_names, prod)).lower()
        return any(p in name for p in pf)

    lines = []
    for o in orders:
        status = int(o.get("status") or 0)
        agent = lookup(agent_names, o.get("agent")) or ref_key(o.get("agent"))
        cl = o.get("client") or {}
        client_name = cl.get("clientName") or lookup(client_names, cl) or ref_key(cl)
        for p in o.get("orderProducts") or []:
            prod = p.get("product") or {}
            if not product_ok(prod):
                continue
            qty, returned = num(p.get("quantity")), num(p.get("returned"))
            lines.append({
                "date": (o.get("dateDocument") or o.get("dateCreate") or "")[:10],
                "order": o.get("SD_id") or o.get("CS_id"),
                "status": status,
                "status_name": STATUS_NAMES.get(status, str(status)),
                "agent_key": ref_key(o.get("agent")),
                "agent": agent,
                "client_key": ref_key(cl),
                "client": client_name,
                "client_legal": cl.get("clientLegalName") or "",
                "product_key": ref_key(prod),
                "product": prod.get("name") or lookup(product_names, prod) or ref_key(prod),
                "qty": qty,
                "returned": returned,
                "block": num(p.get("block")),
                "price": num(p.get("price")),
                "summa": num(p.get("summa")),
            })

    return_lines = []
    for d in defects:
        status = int(d.get("status") or 0)
        if status == 5:
            continue
        cl = d.get("client") or {}
        for p in d.get("defectProducts") or []:
            prod = p.get("product") or {}
            if not product_ok(prod):
                continue
            return_lines.append({
                "date": (d.get("date") or "")[:10],
                "doc": d.get("SD_id") or d.get("CS_id"),
                "status_name": STATUS_NAMES.get(status, str(status)),
                "agent_key": ref_key(d.get("agent")),
                "agent": lookup(agent_names, d.get("agent")) or ref_key(d.get("agent")),
                "client": lookup(client_names, cl) or ref_key(cl),
                "product": lookup(product_names, prod) or ref_key(prod),
                "qty": num(p.get("quantity")),
                "summa": num(p.get("summa")),
                "reason": d.get("reason") or "",
            })
    return lines, return_lines


def aggregate(lines, return_lines):
    live = [l for l in lines if l["status"] != 5]

    def bucket():
        return {"orders": set(), "clients": set(), "qty_done": 0.0, "block_done": 0.0,
                "sum_done": 0.0, "qty_pending": 0.0, "sum_pending": 0.0, "returned": 0.0,
                "agents": set()}

    by_agent, by_client, by_product = defaultdict(bucket), defaultdict(bucket), defaultdict(bucket)
    for l in live:
        for b in (by_agent[(l["agent_key"], l["agent"])],
                  by_client[(l["client_key"], l["client"], l["client_legal"])],
                  by_product[(l["product_key"], l["product"])]):
            b["orders"].add(l["order"])
            b["agents"].add(l["agent"])
            if l["status"] in DONE_STATUSES:
                b["clients"].add(l["client_key"])
                b["returned"] += l["returned"]
                b["qty_done"] += l["qty"]
                b["block_done"] += l["block"]
                b["sum_done"] += l["summa"]
            elif l["status"] in PENDING_STATUSES:
                b["qty_pending"] += l["qty"]
                b["sum_pending"] += l["summa"]

    ret_by_agent = defaultdict(lambda: [0.0, 0.0])
    for r in return_lines:
        ret_by_agent[r["agent_key"]][0] += r["qty"]
        ret_by_agent[r["agent_key"]][1] += r["summa"]

    return live, by_agent, by_client, by_product, ret_by_agent


def write_xlsx(path, meta, lines, return_lines):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    live, by_agent, by_client, by_product, ret_by_agent = aggregate(lines, return_lines)
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F2937")

    def sheet(title, headers, rows, money_cols=(), first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(list(r))
        for i, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 45)
            if i in money_cols:
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    row[0].number_format = "#,##0"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    done = [l for l in live if l["status"] in DONE_STATUSES]
    pending = [l for l in live if l["status"] in PENDING_STATUSES]
    summary = [
        ("Davr", f"{meta['from']} — {meta['to']}"),
        ("Mahsulot filtri", ", ".join(meta["product_filter"]) or "Hammasi"),
        ("Yuklab olingan vaqt", meta["generated"]),
        ("", ""),
        ("Sotilgan (yetkazilgan + yopilgan), so'm", round(sum(l["summa"] for l in done))),
        ("Sotilgan, dona", round(sum(l["qty"] for l in done), 2)),
        ("Buyurtma ichida qaytarilgan, dona", round(sum(l["returned"] for l in done), 2)),
        ("Sotilgan, blok", round(sum(l["block"] for l in done), 2)),
        ("Jarayonda (yangi + yuborilgan), so'm", round(sum(l["summa"] for l in pending))),
        ("Buyurtmalar soni (bekorlarsiz)", len({l["order"] for l in live})),
        ("Xarid qilgan do'konlar (AKB, yetkazilganlar bo'yicha)", len({l["client_key"] for l in done})),
        ("Ishlagan agentlar", len({l["agent_key"] for l in live})),
        ("Bekor qilingan buyurtmalar", len({l["order"] for l in lines if l["status"] == 5})),
        ("Qaytarish hujjatlari, so'm", round(sum(r["summa"] for r in return_lines))),
    ]
    ws = sheet("Xulosa", ["Ko'rsatkich", "Qiymat"], summary, first=True)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        if isinstance(row[0].value, (int, float)):
            row[0].number_format = "#,##0.##"

    agent_rows = sorted(
        ((name, len(b["orders"]), len(b["clients"]), round(b["qty_done"], 2), round(b["block_done"], 2),
          round(b["sum_done"]), round(b["sum_pending"]),
          round(ret_by_agent[key][1]), round(b["sum_done"] - ret_by_agent[key][1]))
         for (key, name), b in by_agent.items()),
        key=lambda r: -r[5],
    )
    sheet("Agentlar",
          ["Agent", "Buyurtmalar", "Do'konlar (AKB)", "Sotildi, dona", "Sotildi, blok",
           "Sotildi, so'm", "Jarayonda, so'm", "Qaytarish, so'm", "Sof sotuv, so'm"],
          agent_rows, money_cols=(6, 7, 8, 9))

    client_rows = sorted(
        ((name, legal, ", ".join(sorted(a for a in b["agents"] if a)), len(b["orders"]),
          round(b["qty_done"], 2), round(b["block_done"], 2), round(b["sum_done"]), round(b["sum_pending"]))
         for (_, name, legal), b in by_client.items()),
        key=lambda r: -r[6],
    )
    sheet("Do'konlar",
          ["Do'kon", "Yuridik nomi", "Agent(lar)", "Buyurtmalar", "Sotildi, dona", "Sotildi, blok",
           "Sotildi, so'm", "Jarayonda, so'm"],
          client_rows, money_cols=(7, 8))

    product_rows = sorted(
        ((name, len(b["clients"]), round(b["qty_done"], 2), round(b["block_done"], 2),
          round(b["sum_done"]), round(b["returned"], 2))
         for (_, name), b in by_product.items()),
        key=lambda r: -r[4],
    )
    sheet("Mahsulotlar",
          ["Mahsulot", "Do'konlar", "Sotildi, dona", "Sotildi, blok", "Sotildi, so'm", "Qaytarilgan, dona"],
          product_rows, money_cols=(5,))

    sheet("Buyurtma qatorlari",
          ["Sana", "Buyurtma", "Holat", "Agent", "Do'kon", "Mahsulot", "Soni", "Qaytarilgan",
           "Blok", "Narx", "Summa"],
          [(l["date"], l["order"], l["status_name"], l["agent"], l["client"], l["product"], l["qty"],
            l["returned"], l["block"], l["price"], l["summa"]) for l in lines],
          money_cols=(10, 11))

    sheet("Qaytarishlar",
          ["Sana", "Hujjat", "Holat", "Agent", "Do'kon", "Mahsulot", "Soni", "Summa", "Sabab"],
          [(r["date"], r["doc"], r["status_name"], r["agent"], r["client"], r["product"], r["qty"],
            r["summa"], r["reason"]) for r in return_lines],
          money_cols=(8,))

    wb.save(path)


def main():
    today = dt.date.today()
    ap = argparse.ArgumentParser(description="SalesDoc oylik sotuv hisobotini Excel'ga yuklash")
    ap.add_argument("--from", dest="date_from", default=today.replace(day=1).isoformat())
    ap.add_argument("--to", dest="date_to", default=today.isoformat())
    ap.add_argument("--product", action="append", default=[],
                    help="Mahsulot nomida shu so'z bo'lsa olinadi (bir necha marta berish mumkin)")
    ap.add_argument("--out", default=None, help="Excel fayl yo'li")
    ap.add_argument("--raw", default=None, help="Xom JSON'ni ham saqlash uchun fayl yo'li")
    args = ap.parse_args()

    missing = [v for v in ("SALESDOC_DOMAIN", "SALESDOC_LOGIN", "SALESDOC_PASSWORD") if not os.environ.get(v)]
    if missing:
        sys.exit("Muhit o'zgaruvchilari yo'q: " + ", ".join(missing))

    client = SalesDocClient(
        os.environ["SALESDOC_DOMAIN"], os.environ["SALESDOC_LOGIN"], os.environ["SALESDOC_PASSWORD"],
        os.environ.get("SALESDOC_FILIAL") or None,
    )
    client.login()
    agents, clients, products, orders, defects = fetch(client, args.date_from, args.date_to)
    if args.raw:
        with open(args.raw, "w", encoding="utf-8") as f:
            json.dump({"agents": agents, "clients": clients, "products": products,
                       "orders": orders, "defects": defects}, f, ensure_ascii=False)

    lines, return_lines = build_report(agents, clients, products, orders, defects, args.product)
    out = args.out or f"salesdoc_{args.date_from}_{args.date_to}.xlsx"
    meta = {"from": args.date_from, "to": args.date_to, "product_filter": args.product,
            "generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M")}
    write_xlsx(out, meta, lines, return_lines)
    print(f"Tayyor: {out}  (buyurtmalar: {len(orders)}, qatorlar: {len(lines)}, "
          f"qaytarishlar: {len(return_lines)})", file=sys.stderr)


if __name__ == "__main__":
    main()
